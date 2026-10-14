"""Data models for Watchlist Companies, Sources, and Analysts, plus Excel sync."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import openpyxl

from .config import CONTROL_PANEL_PATH


@dataclass
class Company:
    id: str
    name: str
    nse_symbol: str = ""
    bse_code: str = ""
    aliases: list[str] = field(default_factory=list)
    search_queries: list[str] = field(default_factory=list)
    analyst_names: list[str] = field(default_factory=list)
    folder: str = ""
    folders: list[str] = field(default_factory=list)
    is_active: bool = True

    @property
    def all_names(self) -> list[str]:
        """All names and aliases to match against in headlines/articles."""
        names = [self.name]
        for a in self.aliases:
            if a and a not in names:
                names.append(a)
        return names


@dataclass
class Source:
    name: str
    url: str
    type: str = "news"          # filing | news | rating | press | page | sitemap | rss
    company: str | None = None  # if page/press belongs to a specific company
    keep_pattern: str | None = None
    language: str = "en"
    body_check: bool = False
    is_active: bool = True


@dataclass
class Analyst:
    id: str
    name: str
    phone: str
    email: str = ""
    assigned_companies: list[str] = field(default_factory=list)
    assigned_folders: list[str] = field(default_factory=list)


def create_starter_excel(file_path: Path = CONTROL_PANEL_PATH) -> None:
    """Create a structured starter Excel file with sample analysts and companies."""
    wb = openpyxl.Workbook()

    # Sheet 1: Watchlist
    ws_watch = wb.active
    if ws_watch is None:
        ws_watch = wb.create_sheet(title="Watchlist")
    else:
        ws_watch.title = "Watchlist"
    ws_watch.append(["Company Name", "NSE Symbol", "BSE Code", "Aliases (comma separated)", "Search Queries (comma separated)", "Assigned Analysts", "Active"])
    
    sample_companies = [
        ("Tata Motors Ltd", "TATAMOTORS", "500570", "Tata Motors, JLR, Jaguar Land Rover, TaMo", '"Tata Motors" OR "JLR"', "Rahul", "YES"),
        ("Polycab India Ltd", "POLYCAB", "542652", "Polycab, Polycab Wires, Polycab India", '"Polycab" OR "Polycab India"', "Rahul", "YES"),
        ("Reliance Industries Ltd", "RELIANCE", "500325", "Reliance, RIL, Jio, Reliance Retail", '"Reliance Industries" OR "RIL"', "Priya", "YES"),
        ("Infosys Ltd", "INFY", "500209", "Infosys, Infy, Salil Parekh", '"Infosys" OR "Infy"', "Priya", "YES"),
        ("HDFC Bank Ltd", "HDFCBANK", "500180", "HDFC Bank, HDFC, Sashidhar Jagdishan", '"HDFC Bank"', "Amit", "YES"),
    ]
    for row in sample_companies:
        ws_watch.append(row)

    # Sheet 2: Analysts
    ws_analysts = wb.create_sheet(title="Analysts")
    ws_analysts.append(["Analyst Name", "WhatsApp Phone", "Email", "Assigned Companies (comma separated)"])
    sample_analysts = [
        ("Rahul", "9820000001", "rahul@vriddhicapital.com", "Tata Motors Ltd, Polycab India Ltd"),
        ("Priya", "9820000002", "priya@vriddhicapital.com", "Reliance Industries Ltd, Infosys Ltd"),
        ("Amit", "9820000003", "amit@vriddhicapital.com", "HDFC Bank Ltd"),
    ]
    for row in sample_analysts:
        ws_analysts.append(row)

    # Sheet 3: Sources
    ws_sources = wb.create_sheet(title="Sources")
    ws_sources.append(["Source Name", "Type", "URL", "Keep URL Pattern", "Body Check", "Active"])
    sample_sources = [
        ("Livemint News Sitemap", "sitemap", "https://www.livemint.com/sitemap/news.xml", r"/(companies|market)/", "YES", "YES"),
        ("Economic Times RSS", "rss", "https://economictimes.indiatimes.com/markets/stocks/rssfeeds/2146842.cms", "", "YES", "YES"),
        ("Moneycontrol RSS", "rss", "https://www.moneycontrol.com/rss/MCtopnews.xml", "", "NO", "YES"),
        ("BSE Corporate Announcements", "filing", "https://www.bseindia.com/corporates/ann.html", "", "NO", "YES"),
        ("NSE Corporate Announcements", "filing", "https://www.nseindia.com/companies-listing/corporate-filings-announcements", "", "NO", "YES"),
    ]
    for row in sample_sources:
        ws_sources.append(row)

    wb.save(file_path)
